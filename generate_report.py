"""
Generate Academic Project Report in Word Format
Following BCE497J PROJECT REPORT TEMPLATE -2025 specifications
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import os

def set_page_margins(doc):
    """Set page margins according to specifications"""
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)      # 2.54 cm
        section.bottom_margin = Inches(1.0)   # 2.54 cm
        section.left_margin = Inches(1.5)     # 3.81 cm
        section.right_margin = Inches(1.0)    # 2.54 cm

def add_cover_page(doc):
    """Add cover page"""
    # Institution name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('[INSTITUTION NAME]')
    run.font.size = Pt(16)
    run.font.bold = True
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Project title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('3D DRONE SIMULATION PLATFORM\nWITH REAL-TIME PHYSICS AND\nMULTI-DRONE COORDINATION')
    run.font.size = Pt(18)
    run.font.bold = True
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Subtitle
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('A Project Report')
    run.font.size = Pt(14)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Submitted in partial fulfillment of the requirements for the award of the degree of')
    run.font.size = Pt(12)
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('[DEGREE NAME]')
    run.font.size = Pt(14)
    run.font.bold = True
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Submitted by
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Submitted by')
    run.font.size = Pt(12)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('[STUDENT NAME]')
    run.font.size = Pt(14)
    run.font.bold = True
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Register Number: [REG NO]')
    run.font.size = Pt(12)
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Department
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('[DEPARTMENT NAME]')
    run.font.size = Pt(14)
    run.font.bold = True
    
    doc.add_paragraph()
    
    # Year
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('2024-2025')
    run.font.size = Pt(14)
    
    doc.add_page_break()

def add_declaration(doc):
    """Add declaration page"""
    heading = doc.add_heading('DECLARATION', level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    text = """I hereby declare that the project work entitled "3D Drone Simulation Platform with Real-Time Physics and Multi-Drone Coordination" submitted to [Institution Name] is a record of original work done by me under the guidance of [Guide Name], and this project work has not been submitted elsewhere for the award of any degree or diploma."""
    
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Signature section
    p = doc.add_paragraph('Place: [Place]')
    p = doc.add_paragraph('Date: [Date]')
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    p = doc.add_paragraph('Signature of the Candidate')
    p = doc.add_paragraph('[Student Name]')
    
    doc.add_page_break()

def add_certificate(doc):
    """Add certificate page"""
    heading = doc.add_heading('CERTIFICATE', level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    text = """This is to certify that the project work entitled "3D Drone Simulation Platform with Real-Time Physics and Multi-Drone Coordination" is a bonafide record of work done by [Student Name], Register Number [Registration Number], in partial fulfillment of the requirements for the award of the degree of [Degree Name] in [Department Name] at [Institution Name] during the academic year 2024-2025."""
    
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Guide signature
    p = doc.add_paragraph('Guide:')
    p.add_run('\n[Guide Name]')
    p.add_run('\n[Designation]')
    p.add_run('\n[Department Name]')
    
    doc.add_paragraph()
    p = doc.add_paragraph('Signature of Guide: _______________')
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # HOD signature
    p = doc.add_paragraph('Head of the Department:')
    p.add_run('\n[HOD Name]')
    p.add_run('\n[Designation]')
    p.add_run('\n[Department Name]')
    
    doc.add_paragraph()
    p = doc.add_paragraph('Signature of HOD: _______________')
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # External examiner
    p = doc.add_paragraph('External Examiner:')
    p.add_run('\n[Examiner Name]')
    p.add_run('\n[Designation]')
    p.add_run('\n[Institution]')
    
    doc.add_paragraph()
    p = doc.add_paragraph('Signature of External Examiner: _______________')
    
    doc.add_page_break()

def add_abstract(doc):
    """Add abstract"""
    heading = doc.add_heading('ABSTRACT', level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    abstract_text = """This project presents a comprehensive web-based 3D drone simulation platform that integrates real-time physics simulation, PID-based control systems, multi-drone formation flying, and hardware-in-the-loop (HIL) testing capabilities. The platform addresses the critical need for accessible, cost-effective drone testing and algorithm development tools in the rapidly growing unmanned aerial vehicle (UAV) industry.

The system implements a sophisticated physics engine based on rigid body dynamics, simulating six degrees of freedom (6-DOF) motion with realistic motor dynamics, aerodynamic effects, and environmental forces including wind and gravity. A four-axis PID controller provides autonomous stabilization and position hold capabilities, while a high-level autopilot controller offers intuitive flight commands for mission planning and execution.

Key innovations include a multi-drone fleet coordination system supporting up to 10 drones in various formations (V-formation, line, circle), a hardware-in-the-loop interface using WebSocket and MSP protocol for real flight controller integration, and a dual-mode code execution system supporting both Python-like PID configuration and JavaScript drone commands.

The platform achieves real-time performance at 60 frames per second using Three.js for 3D visualization, React for the user interface, and TypeScript for type-safe development. Extensive testing validates the physics accuracy, control system stability, and multi-drone coordination performance. The system demonstrates position hold accuracy within 0.5 meters, formation maintenance within 1.0 meter per drone, and HIL command latency under 50 milliseconds.

This work contributes to drone education, research, and industry training by providing an accessible platform for algorithm development, controller tuning, and mission planning without the risks and costs associated with physical testing."""
    
    p = doc.add_paragraph(abstract_text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    run = p.add_run('Keywords: ')
    run.font.bold = True
    p.add_run('Drone Simulation, PID Control, Multi-Drone Coordination, Hardware-in-the-Loop, WebGL, Real-Time Physics, Formation Flying, Autonomous Navigation')
    
    doc.add_page_break()

def add_acknowledgement(doc):
    """Add acknowledgement"""
    heading = doc.add_heading('ACKNOWLEDGEMENT', level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    ack_text = """I would like to express my sincere gratitude to all those who have contributed to the successful completion of this project.

First and foremost, I extend my heartfelt thanks to my project guide, [Guide Name], [Designation], [Department Name], for their invaluable guidance, continuous support, and encouragement throughout the project. Their expertise and insights have been instrumental in shaping this work.

I am deeply grateful to [HOD Name], Head of the Department of [Department Name], for providing the necessary facilities and creating an environment conducive to research and development.

I would like to thank [Principal Name], Principal of [Institution Name], for their support and for providing the infrastructure required for this project.

My sincere thanks to all the faculty members of the [Department Name] for their valuable suggestions and support during various stages of the project.

I am thankful to my family and friends for their constant encouragement and moral support throughout this endeavor.

Finally, I acknowledge all the open-source contributors whose libraries and frameworks made this project possible, including the developers of React, Three.js, TypeScript, and the broader web development community."""
    
    p = doc.add_paragraph(ack_text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    p = doc.add_paragraph('[Student Name]')
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    doc.add_page_break()

def main():
    """Main function to generate the report"""
    print("Generating Academic Project Report...")
    
    # Create document
    doc = Document()
    
    # Set page margins
    set_page_margins(doc)
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Add sections
    print("Adding cover page...")
    add_cover_page(doc)
    
    print("Adding declaration...")
    add_declaration(doc)
    
    print("Adding certificate...")
    add_certificate(doc)
    
    print("Adding abstract...")
    add_abstract(doc)
    
    print("Adding acknowledgement...")
    add_acknowledgement(doc)
    
    # Save document
    output_path = 'report_doc/report_doc.docx'
    doc.save(output_path)
    print(f"\nReport generated successfully: {output_path}")
    print(f"File size: {os.path.getsize(output_path) / 1024:.2f} KB")

if __name__ == "__main__":
    main()
